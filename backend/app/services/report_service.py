"""
Report Service (MongoDB IRS Database)
Handles report generation for trips directly into trip_reports
Includes comprehensive school visit history (lịch sử đi qua các điểm trường)
"""

import logging
import os
import io
import asyncio
from typing import List, Optional, Dict
from datetime import datetime, timedelta, timezone
import uuid
import re
import httpx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.database import get_database
from app.schemas import WaypointType

logger = logging.getLogger(__name__)

REPORT_API_URL = os.getenv("REPORT_API_URL", "http://localhost:8001/report")
REPORT_API_TIMEOUT = float(os.getenv("REPORT_API_TIMEOUT", "2.0"))

_memory_job_status: Dict[str, dict] = {}


def to_vietnam_time(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, str):
        clean_str = dt.strip()
        if re.match(r"^\d{2}/\d{2}/\d{4}", clean_str):
            return clean_str
        try:
            dt = datetime.fromisoformat(clean_str.replace("Z", "+00:00"))
        except Exception:
            try:
                dt = datetime.strptime(clean_str[:19], "%Y-%m-%d %H:%M:%S")
            except Exception:
                return dt
    if hasattr(dt, "tzinfo"):
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        vn_time = dt.astimezone(timezone(timedelta(hours=7)))
        return vn_time.strftime('%d/%m/%Y %H:%M')
    return str(dt)


def to_vietnam_date(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, str):
        clean_str = dt.strip()
        if re.match(r"^\d{2}/\d{2}/\d{4}", clean_str):
            return clean_str
        try:
            dt = datetime.fromisoformat(clean_str.replace("Z", "+00:00"))
        except Exception:
            try:
                dt = datetime.strptime(clean_str[:10], "%Y-%m-%d")
            except Exception:
                return dt
    if hasattr(dt, "tzinfo"):
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        vn_time = dt.astimezone(timezone(timedelta(hours=7)))
        return vn_time.strftime('%d/%m/%Y')
    return str(dt)


def extract_trip_points(trip: dict):
    if "schools_data" in trip and "rest_stops_data" in trip:
        return trip.get("schools_data") or [], trip.get("rest_stops_data") or []
    # Ưu tiên waypoints (dữ liệu động có tickets, visit_logs) trước destinations (mảng tĩnh ban đầu)
    raw_points = trip.get("waypoints") or trip.get("schools") or trip.get("destinations") or []
    schools = []
    rest_stops = []
    for p in raw_points:
        p_type = str(p.get("type") or "").upper()
        if p_type in ["REST_STOP", WaypointType.REST_STOP.value.upper()]:
            rest_stops.append(p)
        else:
            schools.append(p)
    if not rest_stops and trip.get("stops"):
        rest_stops = trip.get("stops", [])
    return schools, rest_stops


def get_point_tickets(p: dict) -> int:
    t_list = p.get("tickets") or []
    if t_list:
        valid_tickets = [t for t in t_list if not t.get("is_deleted")]
        collected_sum = sum(int(t.get("tickets_collected", 0)) for t in valid_tickets)
        if collected_sum > 0:
            return collected_sum
        if len(valid_tickets) > 0:
            return len(valid_tickets)
    try:
        if p.get("tickets_collected") is not None:
            return int(p["tickets_collected"])
        if p.get("total_tickets") is not None:
            return int(p["total_tickets"])
    except (ValueError, TypeError):
        pass
    return 0


def is_point_visited(p: dict) -> bool:
    return bool(p.get("is_visited") or p.get("visited") or False)


async def build_report_prompt(trip: dict, created_at: Optional[datetime] = None) -> str:
    schools, rest_stops = extract_trip_points(trip)
    schools_visited = [w for w in schools if is_point_visited(w)]
    schools_unvisited = [w for w in schools if not is_point_visited(w)]
    total_tickets = sum(get_point_tickets(w) for w in schools)
    report_time = created_at or datetime.now(timezone.utc)

    prompt = f"""# Tạo Báo Cáo Chuyến Đi Tuyển Sinh

Vui lòng tạo một báo cáo chi tiết và chuyên nghiệp cho chuyến đi tuyển sinh sau đây. Báo cáo cần được viết bằng tiếng Việt, sử dụng định dạng Markdown và bao gồm các phần sau:

## Thông Tin Chuyến Đi
- **Tên chuyến đi**: {trip.get('name')}
- **Ngày lập báo cáo**: {to_vietnam_time(report_time)}
- **Trạng thái**: {trip.get('status')}
- **Ngày tạo chuyến đi**: {to_vietnam_time(trip.get('created_at'))}
- **Tổng số trường**: {len(schools)} trường
- **Trường đã đi**: {len(schools_visited)}/{len(schools)} trường
- **Tổng số phiếu thu được**: {total_tickets} phiếu

## Danh Sách Trường Học

### Các trường đã đến ({len(schools_visited)} trường)
"""

    if schools_visited:
        for idx, school in enumerate(schools_visited, 1):
            prompt += f"\n#### {idx}. {school.get('name')}\n"
            prompt += f"- **Địa chỉ**: {school.get('address') or 'Chưa có thông tin'}\n"
            prompt += f"- **Ngày đến**: {to_vietnam_time(school.get('visited_at'))}\n"
            st_count = get_point_tickets(school)
            prompt += f"- **Số phiếu thu được**: {st_count} phiếu\n"
            v_count = school.get("total_visits") or 1
            prompt += f"- **Lịch sử ghé thăm**: Đã ghé thăm {v_count} lần\n"
            if school.get("principal_name") or school.get("vice_principal_name"):
                prompt += f"- **Ban giám hiệu**:\n"
                if school.get("principal_name"):
                    prompt += f"  - Hiệu trưởng: {school.get('principal_name')}"
                    if school.get("principal_phone"):
                        prompt += f" (SĐT: {school.get('principal_phone')})"
                    prompt += "\n"
                if school.get("vice_principal_name"):
                    prompt += f"  - Phó hiệu trưởng: {school.get('vice_principal_name')}"
                    if school.get("vice_principal_phone"):
                        prompt += f" (SĐT: {school.get('vice_principal_phone')})"
                    prompt += "\n"
            if school.get("notes"):
                prompt += f"- **Ghi chú**: {school.get('notes')}\n"
    else:
        prompt += "\n*Chưa có trường nào được ghé thăm.*\n"

    prompt += f"\n### Các trường chưa đến ({len(schools_unvisited)} trường)\n\n"
    if schools_unvisited:
        for idx, school in enumerate(schools_unvisited, 1):
            prompt += f"{idx}. **{school.get('name')}**"
            if school.get("address"):
                prompt += f" - {school.get('address')}"
            prompt += "\n"

    return prompt


def build_native_admission_report(trip: dict, created_at: Optional[datetime] = None) -> str:
    schools, _ = extract_trip_points(trip)
    schools_visited = [w for w in schools if is_point_visited(w)]

    total_tickets = sum(get_point_tickets(w) for w in schools)
    if not total_tickets and trip.get("total_tickets"):
        total_tickets = trip.get("total_tickets", 0)

    report_time = created_at or datetime.now(timezone.utc)

    lines = []
    lines.append("# BÁO CÁO CÔNG TÁC TUYỂN SINH THỰC ĐỊA")
    lines.append("")
    lines.append("## 1. Thông Tin Chuyến Đi")
    lines.append(f"- **Tên chuyến đi**: {trip.get('name') or 'Lộ trình tuyển sinh'}")
    lines.append(f"- **Mã chuyến đi**: `{trip.get('trip_code') or trip.get('id', '')}`")
    if trip.get("campaign_id"):
        lines.append(f"- **Mã chiến dịch liên kết**: `{trip.get('campaign_id')}`")
    lines.append(f"- **Thời điểm lập báo cáo**: {to_vietnam_time(report_time)}")
    lines.append(f"- **Trạng thái chuyến đi**: {trip.get('status', 'Hoàn thành')}")
    if trip.get("created_at"):
        lines.append(f"- **Thời gian khởi tạo**: {to_vietnam_time(trip.get('created_at'))}")
    
    start_pt = trip.get("start_point") or {}
    staff_list = start_pt.get("assigned_staff") or trip.get("assigned_staff") or []
    if staff_list:
        staff_names = [s.get("name") for s in staff_list if s.get("name")]
        lines.append(f"- **Đoàn cán bộ tuyển sinh**: {', '.join(staff_names)}")
    elif trip.get("team", {}).get("leader_name"):
        lines.append(f"- **Trưởng đoàn**: {trip['team']['leader_name']}")
    lines.append("")

    lines.append("## 2. Kết Quả Nhanh")
    lines.append("")
    lines.append(f"- **Số trường đã tiếp cận**: **{len(schools_visited)}** trường (Tổng số trường mục tiêu: {len(schools)} trường)")
    lines.append(f"- **Tổng số phiếu tuyển sinh thu được**: **{total_tickets}** phiếu")
    lines.append("")

    lines.append("## 3. Bảng Kết Quả Thực Địa Các Điểm Trường")
    lines.append("")
    lines.append("| STT | Tên Trường Học | Địa Chỉ | Thời Gian Ghé | Ban Giám Hiệu / Người Tiếp Đón | Số Phiếu Thu | Ghi Chú / Đánh Giá Thực Địa |")
    lines.append("| :---: | :--- | :--- | :---: | :--- | :---: | :--- |")
    
    for idx, school in enumerate(schools, 1):
        is_v = is_point_visited(school)
        v_time = to_vietnam_time(school.get("visited_at")) if is_v else "Chưa tiếp cận"
        s_tickets = get_point_tickets(school)
        
        contacts = []
        if school.get("principal_name"):
            contacts.append(f"HT: {school.get('principal_name')}" + (f" ({school.get('principal_phone')})" if school.get('principal_phone') else ""))
        if school.get("vice_principal_name"):
            contacts.append(f"Phó HT: {school.get('vice_principal_name')}" + (f" ({school.get('vice_principal_phone')})" if school.get('vice_principal_phone') else ""))
        if school.get("representative_name"):
            contacts.append(f"Đại diện: {school.get('representative_name')}")
        contact_str = "; ".join(contacts) if contacts else (school.get("contact_name") or school.get("phone") or "Chưa cập nhật")
        
        addr = school.get("address") or "Chưa có địa chỉ"
        note = school.get("notes") or ("Đã tiếp cận và làm việc trực tiếp" if is_v else "Chưa tiếp cận")
        lines.append(f"| {idx} | **{school.get('name')}** | {addr} | {v_time} | {contact_str} | **{s_tickets}** | {note} |")
    lines.append("")

    lines.append("## 4. Nhận Xét & Đề Xuất Của Đoàn")
    lines.append("")
    lines.append("- **Thuận lợi & Đánh giá chung**: Ban Giám hiệu các trường nhiệt tình phối hợp; học sinh chủ động tìm hiểu các chuyên ngành và chỉ tiêu tuyển sinh.")
    lines.append(f"- **Kết quả thu phiếu**: Đạt tổng cộng **{total_tickets}** phiếu khảo sát nguyện vọng.")
    lines.append("- **Đề xuất cho các đợt tiếp theo**: Phân loại phiếu thu theo nguyện vọng ngành học và chuyển giao ngay cho bộ phận tư vấn tuyển sinh gọi điện hướng dẫn hồ sơ.")
    lines.append("")

    return "\n".join(lines)


async def call_report_api(prompt: str, trip: Optional[dict] = None, created_at: Optional[datetime] = None) -> str:
    try:
        async with httpx.AsyncClient(timeout=REPORT_API_TIMEOUT) as client:
            response = await client.post(
                REPORT_API_URL,
                json={"prompt": prompt, "model": "Qwen/Qwen3-4B-Instruct-2507", "async_mode": False},
                headers={"Content-Type": "application/json"}
            )
            if response.status_code == 200:
                data = response.json()
                content = data.get("result") or data.get("report_content")
                if content and len(content) > 100:
                    return content
    except Exception as e:
        logger.info(f"External report LLM not available ({e}), using instant native Vietnamese admission report.")
    
    return build_native_admission_report(trip, created_at)


def generate_report_excel_bytes(report: dict, trip: Optional[dict] = None) -> bytes:
    wb = openpyxl.Workbook()

    font_title = Font(name="Calibri", size=15, bold=True, color="1F497D")
    font_section = Font(name="Calibri", size=12, bold=True, color="1F497D")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_sub_header = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_accent = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    fill_success = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_warning = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    trip_name = report.get("trip_name") or (trip.get("name") if trip else "Lộ trình tuyển sinh")
    created_at_str = to_vietnam_time(report.get("created_at"))
    total_schools = report.get("total_schools", 0)
    schools_visited = report.get("schools_visited", 0)
    total_tickets = report.get("total_tickets", 0)
    completion_rate = f"{round(schools_visited / total_schools * 100, 1)}%" if total_schools > 0 else "0%"

    if report.get("schools_data") is not None:
        schools = report.get("schools_data", [])
        rest_stops = report.get("rest_stops_data") or []
    elif trip:
        schools, rest_stops = extract_trip_points(trip)
    else:
        schools, rest_stops = [], []

    # SHEET 1: TỔNG QUAN
    ws1 = wb.active
    ws1.title = "Tổng quan"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:E1")
    ws1["A1"] = "BÁO CÁO TỔNG KẾT CHIẾN DỊCH TUYỂN SINH"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = align_center

    ws1.append([])
    ws1.append(["THÔNG TIN CHUNG", "", "", "", ""])
    ws1.cell(row=3, column=1).font = font_section

    info_rows = [
        ("Tên chuyến đi / Lộ trình:", trip_name),
        ("Mã chuyến đi (ID):", report.get("trip_id") or ""),
        ("Mã chiến dịch liên kết:", report.get("campaign_id") or "N/A"),
        ("Thời điểm lập báo cáo:", created_at_str),
        ("Trạng thái chuyến đi:", "Hoàn thành")
    ]
    for label, val in info_rows:
        ws1.append([label, val, "", "", ""])
        r = ws1.max_row
        ws1.cell(row=r, column=1).font = font_bold
        ws1.cell(row=r, column=2).font = font_normal

    ws1.append([])
    ws1.append(["CHỈ SỐ THỰC ĐẠT (KPIs)", "", "", "", ""])
    kpi_title_row = ws1.max_row
    ws1.cell(row=kpi_title_row, column=1).font = font_section

    kpi_headers = ["Chỉ số đánh giá", "Giá trị thực đạt", "Đơn vị", "Ghi chú"]
    ws1.append(kpi_headers)
    kpi_head_row = ws1.max_row
    for col_idx, h in enumerate(kpi_headers, 1):
        cell = ws1.cell(row=kpi_head_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    kpi_data = [
        ["Tổng số trường mục tiêu", total_schools, "Trường", "Toàn bộ điểm trường được quy hoạch"],
        ["Số trường đã tiếp cận thành công", schools_visited, "Trường", "Đã check-in và làm việc trực tiếp"],
        ["Tỷ lệ hoàn thành lộ trình", completion_rate, "%", f"{schools_visited}/{total_schools} điểm trường"],
        ["Tổng số phiếu tuyển sinh thu được", total_tickets, "Phiếu", "Tổng hợp từ các đợt phát & thu phiếu"],
        ["Số điểm dừng chân / Trạm nghỉ đã ghé", len([r for r in rest_stops if is_point_visited(r)]), "Điểm", "Hậu cần đoàn công tác"]
    ]
    for row_item in kpi_data:
        ws1.append(row_item)
        r = ws1.max_row
        for col_idx in range(1, 5):
            cell = ws1.cell(row=r, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in (2, 3):
                cell.alignment = align_center
            if col_idx == 2:
                cell.font = font_bold
                cell.fill = fill_accent

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    # SHEET 2: CHI TIẾT TRƯỜNG HỌC (KÈM LỊCH SỬ ĐI QUA)
    ws2 = wb.create_sheet(title="Chi tiết trường học")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:K1")
    ws2["A1"] = f"DANH SÁCH CHI TIẾT & LỊCH SỬ ĐI QUA CÁC ĐIỂM TRƯỜNG - {trip_name}"
    ws2["A1"].font = font_title
    ws2["A1"].alignment = align_center
    ws2.append([])

    school_headers = [
        "STT", "Mã trường", "Tên trường học", "Địa chỉ",
        "Trạng thái", "Lịch sử đi qua (Số lần)", "Thời gian đến", "Số phiếu thu",
        "Nhật ký lịch sử đi qua & Tiếp xúc", "Ban Giám Hiệu / Liên hệ", "Ghi chú thực địa"
    ]
    ws2.append(school_headers)
    head_row = ws2.max_row
    for col_idx in range(1, len(school_headers) + 1):
        cell = ws2.cell(row=head_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    if schools:
        for idx, school in enumerate(schools, 1):
            is_v = is_point_visited(school)
            v_str = "Đã đến" if is_v else "Chưa đến"
            v_time = to_vietnam_time(school.get("visited_at")) if is_v else ""
            t_count = get_point_tickets(school)
            v_count = school.get("total_visits") or (1 if is_v else 0)
            v_count_str = f"{v_count} lần" if v_count > 0 else "Chưa ghé"

            # Tóm tắt lịch sử đi qua
            history_items = []
            for h in (school.get("visit_history") or []):
                h_d = to_vietnam_date(h.get("date") or h.get("visited_at"))
                h_txt = f"{h_d}: {h.get('notes') or h.get('trip_name') or 'Check-in'}"
                history_items.append(h_txt)
            if not history_items and is_v:
                history_items.append(f"{to_vietnam_date(school.get('visited_at')) or 'Đợt này'}: Đã tiếp cận thực địa & thu {t_count} phiếu")
            history_summary = ";\n".join(history_items[:3]) if history_items else "Chưa có lịch sử ghé thăm"

            contacts = []
            if school.get("principal_name"):
                contacts.append(f"HT: {school.get('principal_name')}" + (f" ({school.get('principal_phone')})" if school.get('principal_phone') else ""))
            if school.get("vice_principal_name"):
                contacts.append(f"Phó HT: {school.get('vice_principal_name')}" + (f" ({school.get('vice_principal_phone')})" if school.get('vice_principal_phone') else ""))
            if school.get("representative_name"):
                contacts.append(f"Đại diện: {school.get('representative_name')}")
            contact_str = "; ".join(contacts) if contacts else (school.get("contact_name") or school.get("phone") or "")

            row_vals = [
                idx,
                school.get("school_id") or school.get("id") or "",
                school.get("name") or "",
                school.get("address") or "",
                v_str,
                v_count_str,
                v_time,
                t_count,
                history_summary,
                contact_str,
                school.get("notes") or ""
            ]
            ws2.append(row_vals)
            r = ws2.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws2.cell(row=r, column=col_idx)
                cell.font = font_normal
                cell.border = thin_border
                if col_idx in (1, 5, 6, 7, 8):
                    cell.alignment = align_center
                if col_idx == 5:
                    cell.fill = fill_success if is_v else fill_warning
                if col_idx in (6, 8) and (v_count > 0 or t_count > 0):
                    cell.font = font_bold
    else:
        ws2.append(["1", "", "Chưa có thông tin điểm trường", "", "-", "0 lần", "", 0, "", "", ""])
        ws2.cell(row=ws2.max_row, column=3).font = font_normal

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # SHEET 3: ĐIỂM DỪNG CHÂN
    ws3 = wb.create_sheet(title="Điểm dừng chân")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:F1")
    ws3["A1"] = f"NHẬT KÝ ĐIỂM DỪNG CHÂN / QUÁN ĂN - {trip_name}"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = align_center
    ws3.append([])

    stop_headers = ["STT", "Tên quán / Điểm dừng", "Địa chỉ", "Trạng thái", "Thời gian ghé", "Ghi chú"]
    ws3.append(stop_headers)
    stop_head_row = ws3.max_row
    for col_idx in range(1, len(stop_headers) + 1):
        cell = ws3.cell(row=stop_head_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = thin_border

    if rest_stops:
        for idx, stop in enumerate(rest_stops, 1):
            is_v = is_point_visited(stop)
            v_str = "Đã ghé" if is_v else "Chưa ghé"
            v_time = to_vietnam_time(stop.get("visited_at")) if is_v else ""

            row_vals = [
                idx,
                stop.get("name") or "",
                stop.get("address") or "",
                v_str,
                v_time,
                stop.get("notes") or ""
            ]
            ws3.append(row_vals)
            r = ws3.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws3.cell(row=r, column=col_idx)
                cell.font = font_normal
                cell.border = thin_border
                if col_idx in (1, 4):
                    cell.alignment = align_center
                if col_idx == 4:
                    cell.fill = fill_success if is_v else fill_warning
    else:
        ws3.append(["1", "Không có điểm dừng chân nào trên lộ trình", "", "-", "", ""])
        ws3.cell(row=ws3.max_row, column=2).font = font_normal

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


async def generate_trip_report(trip_id: str, created_at: Optional[datetime] = None) -> dict:
    from app.services.trip_service import trip_service
    db = get_database()
    trip = await trip_service.get_trip(trip_id)
    if not trip:
        raise ValueError(f"Không tìm thấy chuyến đi với ID: {trip_id}")

    schools, rest_stops = extract_trip_points(trip)
    schools_visited = [w for w in schools if is_point_visited(w)]
    total_tickets = sum(get_point_tickets(w) for w in schools)
    if not total_tickets and trip.get("total_tickets"):
        total_tickets = trip.get("total_tickets", 0)

    now = created_at or datetime.now(timezone.utc)

    # Lấy waypoints thực tế từ campaign_waypoints để có tickets, visit_logs, visited_at đầy đủ nhất
    trip_cws = await db.campaign_waypoints.find({"trip_id": trip_id, "is_deleted": {"$ne": True}}).to_list(200)
    cw_by_id = {cw.get("id"): cw for cw in trip_cws if cw.get("id")}
    cw_by_school_id = {cw.get("school_id"): cw for cw in trip_cws if cw.get("school_id")}

    clean_schools = []
    for s in schools:
        s_id = s.get("school_id") or s.get("id")
        s_name = s.get("name")
        cw_match = cw_by_id.get(s.get("id")) or cw_by_school_id.get(s_id)
        
        is_v = is_point_visited(s) or (is_point_visited(cw_match) if cw_match else False)
        s_visited_at = s.get("visited_at") or (cw_match.get("visited_at") if cw_match else None)
        if isinstance(s_visited_at, datetime):
            if s_visited_at.tzinfo is None:
                s_visited_at = s_visited_at.replace(tzinfo=timezone.utc)
            s_visited_at = s_visited_at.isoformat()
        elif isinstance(s_visited_at, str) and s_visited_at and not s_visited_at.endswith("Z") and "+" not in s_visited_at:
            s_visited_at = s_visited_at.replace(" ", "T") + "Z"

        s_tickets = s.get("tickets") or (cw_match.get("tickets") if cw_match else []) or []
        s_logs = s.get("visit_logs") or (cw_match.get("visit_logs") if cw_match else []) or []

        # Tính số phiếu của trường trong chuyến đi này
        tickets_count = get_point_tickets(s)
        if tickets_count == 0 and cw_match:
            tickets_count = get_point_tickets(cw_match)
        if tickets_count == 0 and s_tickets:
            tickets_count = sum(int(t.get("tickets_collected", 0)) for t in s_tickets if not t.get("is_deleted"))

        # 1. Truy vấn thông tin trường trong db.schools để lấy tổng số lần ghé thăm lịch sử
        school_doc = None
        if s_id or s_name:
            or_queries = []
            if s_id:
                or_queries.extend([{"id": s_id}, {"code": s_id}])
            if s_name:
                or_queries.append({"name": s_name})
            try:
                school_doc = await db.schools.find_one({
                    "$or": or_queries,
                    "is_deleted": {"$ne": True}
                })
            except Exception as e:
                logger.warning(f"Error querying school {s_name}: {e}")

        # 2. Truy vấn các chuyến đi khác trong quá khứ đã từng ghé thăm trường này
        past_trips = []
        try:
            past_trips = await db.admission_trips.find({
                "id": {"$ne": trip_id},
                "is_deleted": {"$ne": True},
                "$or": [
                    {"destinations.name": s_name},
                    {"destinations.school_id": s_id},
                    {"waypoints.name": s_name}
                ]
            }).to_list(15)
        except Exception as e:
            logger.warning(f"Error querying past trips for {s_name}: {e}")

        # 3. Tổng hợp danh sách lịch sử ghé thăm (visit_history)
        visit_history = []

        # Lần ghé thăm trong chuyến đi hiện tại (nếu đã ghé)
        if is_v:
            visit_history.append({
                "trip_id": trip_id,
                "trip_name": trip.get("name") or "Chuyến đi này",
                "title": "Chuyến đi này",
                "date": s_visited_at or now,
                "staff": s.get("our_contact_person") or (trip.get("assigned_staff", [{}])[0].get("name") if trip.get("assigned_staff") else None),
                "tickets": tickets_count,
                "notes": s.get("notes") or s.get("contact_process") or "Đã hoàn thành check-in thực địa",
                "is_current": True
            })

        # Các lần ghé thăm trong các chuyến đi trước
        for pt in past_trips:
            p_destinations = pt.get("destinations") or pt.get("waypoints") or []
            for pd in p_destinations:
                if (pd.get("name") == s_name or (s_id and (pd.get("school_id") == s_id or pd.get("id") == s_id))) and is_point_visited(pd):
                    visit_history.append({
                        "trip_id": pt.get("id"),
                        "trip_name": pt.get("name"),
                        "title": f"Chuyến đi: {pt.get('name')}",
                        "date": pd.get("visited_at") or pt.get("created_at"),
                        "staff": pd.get("our_contact_person"),
                        "tickets": get_point_tickets(pd),
                        "notes": pd.get("notes") or "Đã tiếp xúc và trao đổi tuyển sinh đợt trước",
                        "is_current": False
                    })

        # Các visit_logs nhúng sẵn
        for l in s_logs:
            if not l.get("is_deleted"):
                visit_history.append({
                    "trip_id": trip_id,
                    "trip_name": trip.get("name"),
                    "title": "Nhật ký làm việc",
                    "date": l.get("visit_date") or l.get("created_at"),
                    "staff": l.get("created_by_name") or s.get("our_contact_person"),
                    "tickets": l.get("tickets_collected", 0),
                    "notes": l.get("visit_content") or l.get("notes") or "",
                    "is_current": True
                })

        # Sắp xếp lịch sử theo thời gian mới nhất lên đầu
        visit_history = sorted(visit_history, key=lambda x: str(x.get("date") or ""), reverse=True)
        
        # Tính tổng số lần ghé thăm tích lũy
        calc_total_visits = len(visit_history)
        if school_doc and school_doc.get("total_visits", 0) > calc_total_visits:
            calc_total_visits = school_doc.get("total_visits", 0)
        if is_v and calc_total_visits == 0:
            calc_total_visits = 1

        clean_schools.append({
            "id": s.get("id") or (cw_match.get("id") if cw_match else None),
            "school_id": s_id,
            "name": s_name,
            "address": s.get("address") or (school_doc.get("address") if school_doc else ""),
            "is_visited": is_v,
            "visited_at": s_visited_at,
            "total_visits": calc_total_visits,
            "visit_history": visit_history,
            "tickets": s_tickets,
            "tickets_collected": tickets_count,
            "visit_logs": s_logs,
            "principal_name": s.get("principal_name") or (school_doc.get("principal_name") if school_doc else ""),
            "principal_phone": s.get("principal_phone") or (school_doc.get("principal_phone") if school_doc else ""),
            "vice_principal_name": s.get("vice_principal_name") or (school_doc.get("vice_principal_name") if school_doc else ""),
            "vice_principal_phone": s.get("vice_principal_phone") or (school_doc.get("vice_principal_phone") if school_doc else ""),
            "representative_name": s.get("representative_name"),
            "our_contact_person": s.get("our_contact_person"),
            "our_contact_role": s.get("our_contact_role"),
            "contact_process": s.get("contact_process"),
            "notes": s.get("notes")
        })

    # Cập nhật lại total_tickets chính xác theo clean_schools
    calc_total_tickets = sum(cs.get("tickets_collected", 0) for cs in clean_schools)
    if calc_total_tickets > 0:
        total_tickets = calc_total_tickets
    elif not total_tickets and trip.get("total_tickets"):
        total_tickets = trip.get("total_tickets", 0)

    clean_stops = []
    for r in rest_stops:
        clean_stops.append({
            "id": r.get("id"),
            "name": r.get("name"),
            "address": r.get("address"),
            "type": "REST_STOP",
            "is_visited": is_point_visited(r),
            "visited_at": r.get("visited_at"),
            "notes": r.get("notes")
        })

    # Cập nhật trip tạm thời với clean_schools đã làm giàu lịch sử để build report prompt & content
    trip_enriched = dict(trip)
    trip_enriched["schools"] = clean_schools
    trip_enriched["schools_data"] = clean_schools
    trip_enriched["rest_stops_data"] = clean_stops
    trip_enriched["destinations"] = clean_schools + clean_stops
    trip_enriched["total_tickets"] = total_tickets

    prompt = await build_report_prompt(trip_enriched, created_at)
    report_content = await call_report_api(prompt, trip_enriched, created_at)

    report_id = str(uuid.uuid4())
    report_doc = {
        "id": report_id,
        "trip_id": trip_id,
        "trip_name": trip.get("name"),
        "campaign_id": trip.get("campaign_id"),
        "report_content": report_content,
        "total_schools": len(schools),
        "schools_visited": len(schools_visited),
        "total_tickets": total_tickets,
        "schools_data": clean_schools,
        "rest_stops_data": clean_stops,
        "created_at": now
    }

    await db.trip_reports.insert_one(report_doc)

    # Tự động đồng bộ ngược thông tin thăm trường, phiếu thu và BGH vào collection schools & campaign_waypoints
    for cs in clean_schools:
        s_id = cs.get("school_id") or cs.get("id")
        if not s_id:
            continue

        # Cập nhật campaign_waypoints của chuyến đi
        update_cw = {
            "is_visited": cs.get("is_visited", False),
            "visited_at": cs.get("visited_at"),
            "visit_logs": cs.get("visit_logs") or [],
            "notes": cs.get("notes") or ""
        }
        # Chỉ cập nhật tickets nếu có phiếu thực tế, không ghi đè rỗng [] nếu db đã có phiếu
        if cs.get("tickets"):
            update_cw["tickets"] = cs.get("tickets")

        await db.campaign_waypoints.update_one(
            {"trip_id": trip_id, "$or": [{"school_id": s_id}, {"id": cs.get("id")}]},
            {"$set": update_cw}
        )

        # Cập nhật thông tin trường học vào db.schools
        school_updates = {}
        if cs.get("is_visited"):
            school_updates["last_visited_at"] = cs.get("visited_at") or now
        if cs.get("contact_process"):
            school_updates["contact_process"] = cs.get("contact_process")
            school_updates["total_contact_attempts"] = cs.get("total_visits") or 1
        if cs.get("our_contact_person"):
            school_updates["our_contact_person"] = cs.get("our_contact_person")
        if cs.get("our_contact_role"):
            school_updates["our_contact_role"] = cs.get("our_contact_role")
        if cs.get("principal_name"):
            school_updates["principal_name"] = cs.get("principal_name")
            school_updates["school_board.principal_name"] = cs.get("principal_name")
        if cs.get("principal_phone"):
            school_updates["principal_phone"] = cs.get("principal_phone")
            school_updates["school_board.principal_phone"] = cs.get("principal_phone")
        if cs.get("vice_principal_name"):
            school_updates["vice_principal_name"] = cs.get("vice_principal_name")
            school_updates["school_board.vice_principal_name"] = cs.get("vice_principal_name")
        if cs.get("vice_principal_phone"):
            school_updates["vice_principal_phone"] = cs.get("vice_principal_phone")
            school_updates["school_board.vice_principal_phone"] = cs.get("vice_principal_phone")

        if school_updates:
            school_updates["updated_at"] = now
            await db.schools.update_one(
                {"$or": [{"id": s_id}, {"code": s_id}], "is_deleted": {"$ne": True}},
                {"$set": school_updates}
            )

    try:
        from app.services.school_counter_service import sync_all_school_counters
        await sync_all_school_counters()
    except Exception as se:
        logger.warning(f"Error syncing counters: {se}")

    return report_doc


async def get_trip_reports(trip_id: str) -> List[dict]:
    db = get_database()
    cursor = db.trip_reports.find({"trip_id": trip_id, "is_deleted": {"$ne": True}}).sort("created_at", -1)
    reports = await cursor.to_list(length=1000)
    trip = await db.admission_trips.find_one({"id": trip_id})
    trip_name = trip.get("name") if trip else "Chuyến đi"
    for r in reports:
        if not r.get("trip_name"):
            r["trip_name"] = trip_name
    return reports


async def get_all_reports() -> List[dict]:
    db = get_database()
    cursor = db.trip_reports.find({"is_deleted": {"$ne": True}}).sort("created_at", -1)
    reports = await cursor.to_list(length=1000)

    trips_cache = {}
    camps_cache = {}

    report_list = []
    for report in reports:
        t_id = report.get("trip_id")
        if t_id and t_id not in trips_cache:
            trip_doc = await db.admission_trips.find_one({"id": t_id, "is_deleted": {"$ne": True}})
            trips_cache[t_id] = trip_doc
        
        trip = trips_cache.get(t_id)
        if trip:
            if not report.get("trip_name"):
                report["trip_name"] = trip.get("name") or "Chuyến đi"
            report["trip_code"] = trip.get("trip_code") or trip.get("id")
            if not report.get("campaign_id") and trip.get("campaign_id"):
                report["campaign_id"] = trip.get("campaign_id")

            if trip.get("team", {}).get("leader_name"):
                report["leader_name"] = trip["team"]["leader_name"]
            elif trip.get("assigned_staff") and len(trip["assigned_staff"]) > 0:
                report["leader_name"] = trip["assigned_staff"][0].get("name")

        c_id = report.get("campaign_id")
        if c_id and c_id not in camps_cache:
            camp_doc = await db.campaigns.find_one({"$or": [{"id": c_id}, {"_id": c_id}], "is_deleted": {"$ne": True}})
            camps_cache[c_id] = camp_doc

        camp = camps_cache.get(c_id)
        if camp:
            report["campaign_name"] = camp.get("name")
            report["campaign_code"] = camp.get("code") or (f"CAMP-{c_id[:6].upper()}" if c_id else "")
            if not report.get("leader_name") and camp.get("team", {}).get("leader_name"):
                report["leader_name"] = camp["team"]["leader_name"]

        report_list.append(report)

    return report_list


async def get_campaign_summary_report(campaign_id: str) -> dict:
    """
    Tổng hợp báo cáo tiến độ và kết quả toàn chiến dịch theo thời gian thực (Real-time).
    Tự động quét và gom toàn bộ các đợt (admission_trips) hiện có:
    - BẢNG 1: Tổng hợp tiến độ & phiếu thu theo từng đợt
    - BẢNG 2: Tổng hợp kết quả tích lũy theo từng điểm trường
    """
    db = get_database()
    camp = await db.campaigns.find_one({"$or": [{"id": campaign_id}, {"_id": campaign_id}], "is_deleted": {"$ne": True}})
    
    # Lấy toàn bộ chuyến đi thuộc chiến dịch, sắp xếp theo thời gian tạo
    cursor = db.admission_trips.find({"campaign_id": campaign_id, "is_deleted": {"$ne": True}}).sort("created_at", 1)
    trips = await cursor.to_list(length=100)
    
    # Lấy toàn bộ báo cáo của các chuyến đi
    trip_ids = [t["id"] for t in trips]
    rep_cursor = db.trip_reports.find({
        "$or": [
            {"campaign_id": campaign_id},
            {"trip_id": {"$in": trip_ids}}
        ],
        "is_deleted": {"$ne": True}
    }).sort("created_at", 1)
    reports = await rep_cursor.to_list(length=100)
    trip_report_map = {r.get("trip_id"): r.get("id") for r in reports if r.get("trip_id")}
    trip_report_doc_map = {r.get("trip_id"): r for r in reports if r.get("trip_id")}

    rounds_data = []
    schools_map = {}

    for idx, trip in enumerate(trips, 1):
        r_name = f"Đợt {idx}"
        trip_code = trip.get("trip_code") or trip.get("id")
        t_id = trip["id"]
        
        leader = "Chưa phân công"
        if trip.get("team", {}).get("leader_name"):
            leader = trip["team"]["leader_name"]
        elif trip.get("assigned_staff") and len(trip["assigned_staff"]) > 0:
            leader = trip["assigned_staff"][0].get("name", "Cán bộ tuyển sinh")
        elif camp and camp.get("team", {}).get("leader_name"):
            leader = camp["team"]["leader_name"]

        # Truy vấn waypoints thực tế từ campaign_waypoints để có tickets
        cw_cursor = db.campaign_waypoints.find({"trip_id": t_id, "is_deleted": {"$ne": True}})
        trip_cws = await cw_cursor.to_list(length=200)

        trip_enriched_copy = dict(trip)
        if trip_cws:
            trip_enriched_copy["waypoints"] = trip_cws

        schools, _ = extract_trip_points(trip_enriched_copy)
        visited_schools = [s for s in schools if is_point_visited(s)]
        tickets_in_trip = sum(get_point_tickets(s) for s in schools)

        # Ưu tiên lấy từ báo cáo đã lập nếu có
        rep_doc = trip_report_doc_map.get(t_id)
        if rep_doc and rep_doc.get("total_tickets") is not None:
            if rep_doc.get("total_tickets") > tickets_in_trip:
                tickets_in_trip = rep_doc.get("total_tickets", 0)
        elif not tickets_in_trip and trip.get("total_tickets"):
            tickets_in_trip = trip.get("total_tickets", 0)

        t_date = to_vietnam_date(trip.get("start_date") or trip.get("created_at"))

        rounds_data.append({
            "round_number": idx,
            "round_name": r_name,
            "trip_id": t_id,
            "trip_code": trip_code,
            "trip_name": trip.get("name") or f"Chuyến đi {r_name}",
            "date": t_date,
            "leader_name": leader,
            "schools_count": len(visited_schools) if visited_schools else len(schools),
            "tickets_collected": tickets_in_trip,
            "status": "Hoàn thành" if str(trip.get("status", "")).lower() in ["completed", "hoan_thanh", "hoàn thành"] else "Đang triển khai",
            "report_id": trip_report_map.get(t_id)
        })

        for s in schools:
            s_key = s.get("school_id") or s.get("id") or s.get("name")
            s_name = s.get("name") or "Trường học"
            if s_key not in schools_map:
                contacts = []
                if s.get("principal_name"):
                    contacts.append(f"HT: {s.get('principal_name')}" + (f" ({s.get('principal_phone')})" if s.get('principal_phone') else ""))
                if s.get("vice_principal_name"):
                    contacts.append(f"Phó HT: {s.get('vice_principal_name')}" + (f" ({s.get('vice_principal_phone')})" if s.get('vice_principal_phone') else ""))
                if s.get("representative_name"):
                    contacts.append(f"Đại diện: {s.get('representative_name')}")
                contact_str = "; ".join(contacts) if contacts else (s.get("contact_name") or s.get("phone") or "Chưa cập nhật")

                schools_map[s_key] = {
                    "school_id": s_key,
                    "name": s_name,
                    "address": s.get("address") or "Chưa có địa chỉ",
                    "visited_rounds": [],
                    "total_visits": 0,
                    "total_tickets": 0,
                    "contact_person": contact_str,
                    "notes": s.get("notes") or "Điểm trường mục tiêu"
                }

            pt_tickets = get_point_tickets(s)
            has_visited = is_point_visited(s) or pt_tickets > 0
            if has_visited:
                if r_name not in schools_map[s_key]["visited_rounds"]:
                    schools_map[s_key]["visited_rounds"].append(r_name)
                schools_map[s_key]["total_visits"] += 1
                schools_map[s_key]["total_tickets"] += pt_tickets
                if s.get("notes"):
                    schools_map[s_key]["notes"] = s.get("notes")
            elif pt_tickets > 0:
                schools_map[s_key]["total_tickets"] += pt_tickets

    schools_list = list(schools_map.values())
    for s_item in schools_list:
        v_rounds = s_item["visited_rounds"]
        if v_rounds:
            s_item["rounds_display"] = f"{', '.join(v_rounds)} ({len(v_rounds)} lần)"
        else:
            s_item["rounds_display"] = "Chưa ghé"

    total_tickets_all = sum(r["tickets_collected"] for r in rounds_data)

    c_name = camp.get("name", "Chiến dịch tuyển sinh") if camp else "Chiến dịch tuyển sinh"
    c_code = camp.get("code") if camp and camp.get("code") else (f"CAMP-{campaign_id[:6].upper()}" if campaign_id else "CAMP")

    start_d = camp.get("start_date") if camp else None
    end_d = camp.get("end_date") if camp else None

    # Fallback ngày từ các chuyến đi nếu chiến dịch chưa lưu start_date / end_date
    if (not start_d or not end_d) and trips:
        trip_dates = [t.get("start_date") or t.get("created_at") for t in trips if (t.get("start_date") or t.get("created_at"))]
        if trip_dates:
            trip_dates.sort()
            if not start_d:
                start_d = trip_dates[0]
            if not end_d:
                end_d = trip_dates[-1]

    leader_n = ""
    if camp and camp.get("team", {}).get("leader_name"):
        leader_n = camp["team"]["leader_name"]
    elif rounds_data and rounds_data[0].get("leader_name"):
        leader_n = rounds_data[0]["leader_name"]

    # Danh sách cán bộ tuyển sinh toàn chiến dịch
    all_staff_names = []
    if camp and camp.get("team", {}).get("members"):
        for m in camp["team"]["members"]:
            if isinstance(m, dict) and m.get("name") and m["name"] not in all_staff_names:
                all_staff_names.append(m["name"])
    for t in trips:
        t_staff = t.get("assigned_staff") or t.get("team", {}).get("members") or (t.get("start_point", {}).get("assigned_staff") if t.get("start_point") else None) or []
        for s in t_staff:
            if isinstance(s, dict) and s.get("name") and s["name"] not in all_staff_names:
                all_staff_names.append(s["name"])
            elif isinstance(s, str) and s not in all_staff_names:
                all_staff_names.append(s)
    if not all_staff_names and leader_n:
        all_staff_names.append(leader_n)

    summary_doc = {
        "campaign_id": campaign_id,
        "campaign_name": c_name,
        "campaign_code": c_code,
        "leader_name": leader_n,
        "status": camp.get("status", "completed") if camp else "completed",
        "campaign_created_at": camp.get("created_at") if camp else None,
        "start_point": camp.get("start_point") if (camp and camp.get("start_point")) else (trips[0].get("start_point") if trips else None),
        "assigned_staff": all_staff_names,
        "start_date": to_vietnam_date(start_d) if start_d else "",
        "end_date": to_vietnam_date(end_d) if end_d else "",
        "total_rounds": len(rounds_data),
        "total_schools": len(schools_list),
        "total_tickets": total_tickets_all,
        "rounds_data": rounds_data,
        "schools_data": schools_list,
        "trips": [{"id": t["id"], "trip_code": t.get("trip_code", t["id"]), "name": t.get("name"), "report_id": trip_report_map.get(t["id"])} for t in trips],
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    return summary_doc



async def get_report_by_id(report_id: str) -> Optional[dict]:
    db = get_database()
    report = await db.trip_reports.find_one({"id": report_id, "is_deleted": {"$ne": True}})
    if report:
        trip = await db.admission_trips.find_one({"id": report.get("trip_id")})
        if trip:
            if not report.get("trip_name"):
                report["trip_name"] = trip.get("name")
            if not report.get("trip_code"):
                report["trip_code"] = trip.get("trip_code") or trip.get("id")
            if not report.get("campaign_id") and trip.get("campaign_id"):
                report["campaign_id"] = trip.get("campaign_id")
            report["status"] = trip.get("status") or "completed"
            report["trip_status"] = trip.get("status") or "completed"
            report["trip_created_at"] = trip.get("created_at")
            report["start_point"] = trip.get("start_point")
            
            # Cán bộ tuyển sinh
            staff_list = trip.get("assigned_staff") or (trip.get("team", {}).get("members") if trip.get("team") else None) or (trip.get("start_point", {}).get("assigned_staff") if trip.get("start_point") else None) or []
            report["assigned_staff"] = staff_list
            if not report.get("leader_name"):
                if trip.get("team", {}).get("leader_name"):
                    report["leader_name"] = trip["team"]["leader_name"]
                elif staff_list and len(staff_list) > 0 and isinstance(staff_list[0], dict):
                    report["leader_name"] = staff_list[0].get("name")

        # Đảm bảo schools_data có tickets_collected chính xác và visited_at chuẩn ISO UTC
        tot_tickets = report.get("total_tickets") or 0
        schools_data = report.get("schools_data") or []
        if len(schools_data) == 1 and tot_tickets > 0:
            if not schools_data[0].get("tickets_collected"):
                schools_data[0]["tickets_collected"] = tot_tickets
        for sch in schools_data:
            v_at = sch.get("visited_at")
            if v_at and isinstance(v_at, str) and not v_at.endswith("Z") and "+" not in v_at:
                sch["visited_at"] = v_at.replace(" ", "T") + "Z"
            elif v_at and isinstance(v_at, datetime):
                if v_at.tzinfo is None:
                    v_at = v_at.replace(tzinfo=timezone.utc)
                sch["visited_at"] = v_at.isoformat()
    return report


async def start_report_generation_job(trip_id: str, created_at: Optional[datetime] = None) -> dict:
    job_id = str(uuid.uuid4())
    _memory_job_status[job_id] = {
        "id": job_id,
        "trip_id": trip_id,
        "status": "pending",
        "progress": 0,
        "result_report_id": None,
        "error_message": None,
        "created_at": created_at or datetime.now(timezone.utc),
        "updated_at": datetime.now(timezone.utc)
    }
    return _memory_job_status[job_id]


async def process_report_job_background(job_id: str, trip_id: str):
    try:
        if job_id in _memory_job_status:
            _memory_job_status[job_id]["status"] = "processing"
            _memory_job_status[job_id]["progress"] = 25

        job = _memory_job_status.get(job_id, {})
        created_at = job.get("created_at")

        report = await generate_trip_report(trip_id, created_at=created_at)

        try:
            from app.services.school_counter_service import sync_all_school_counters
            await sync_all_school_counters()
        except Exception as se:
            logger.warning(f"Could not re-sync counters after report generation: {se}")

        if job_id in _memory_job_status:
            _memory_job_status[job_id]["status"] = "completed"
            _memory_job_status[job_id]["progress"] = 100
            _memory_job_status[job_id]["result_report_id"] = report["id"]
            _memory_job_status[job_id]["updated_at"] = datetime.now(timezone.utc)
    except Exception as e:
        logger.error(f"Job {job_id} failed: {e}")
        if job_id in _memory_job_status:
            _memory_job_status[job_id]["status"] = "failed"
            _memory_job_status[job_id]["error_message"] = str(e)
            _memory_job_status[job_id]["updated_at"] = datetime.now(timezone.utc)


async def get_report_job_status(job_id: str) -> Optional[dict]:
    if job_id in _memory_job_status:
        return _memory_job_status[job_id]
    db = get_database()
    report = await db.trip_reports.find_one({"id": job_id, "is_deleted": {"$ne": True}})
    if report:
        return {
            "id": job_id,
            "trip_id": report["trip_id"],
            "status": "completed",
            "progress": 100,
            "result_report_id": report["id"],
            "error_message": None
        }
    return None


async def delete_report(report_id: str) -> bool:
    db = get_database()
    now = datetime.now(timezone.utc)
    res = await db.trip_reports.update_one(
        {"id": report_id, "is_deleted": {"$ne": True}},
        {"$set": {"is_deleted": True, "deleted_at": now}}
    )
    return res.modified_count > 0
