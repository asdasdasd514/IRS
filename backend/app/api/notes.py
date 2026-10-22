"""
API Quản lý Ghi chú Thực địa Tuyển sinh (Field Notes Management)
Lưu trữ TRỰC TIẾP bên trong hồ sơ Trường học (collection 'schools' -> mảng 'field_notes')
Thay vì tách bảng riêng biệt, bảo toàn lịch sử thực địa tích lũy theo từng điểm trường đối tác.
"""

import uuid
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, HTTPException, Query, status, Response
from pydantic import BaseModel, Field
from app.core.database import get_database

router = APIRouter(prefix="/notes", tags=["Field Notes"])


class CreateNoteRequest(BaseModel):
    trip_id: Optional[str] = None
    trip_name: Optional[str] = None
    campaign_id: Optional[str] = None
    campaign_name: Optional[str] = None
    school_id: Optional[str] = None
    school_name: Optional[str] = None
    waypoint_id: Optional[str] = None
    category: Optional[str] = "Ghi chú"
    content: str
    author_name: Optional[str] = "Cán bộ tuyển sinh"


class UpdateNoteRequest(BaseModel):
    content: Optional[str] = None
    category: Optional[str] = None
    school_name: Optional[str] = None


async def _get_all_embedded_notes(db) -> List[Dict[str, Any]]:
    """Đọc toàn bộ ghi chú thực địa được nhúng bên trong collection 'schools'"""
    schools_cursor = db.schools.find({"is_deleted": {"$ne": True}})
    all_notes = []

    async for school in schools_cursor:
        s_id = school.get("id") or str(school.get("_id"))
        s_name = school.get("name") or "Trường học"
        fnotes = school.get("field_notes") or []

        # Nếu trường có ghi chú tổng thể mà chưa có trong mảng field_notes
        if not fnotes and school.get("notes") and str(school.get("notes")).strip():
            fnotes = [{
                "id": f"fn-{s_id}",
                "content": str(school["notes"]).strip(),
                "author_name": "Đoàn tuyển sinh",
                "category": "Ghi chú",
                "trip_id": None,
                "trip_name": "Ghi chú trường",
                "campaign_id": None,
                "campaign_name": None,
                "created_at": school.get("updated_at") or school.get("created_at") or datetime.now(timezone.utc),
                "updated_at": school.get("updated_at") or datetime.now(timezone.utc),
                "is_deleted": False
            }]

        for fn in fnotes:
            if not isinstance(fn, dict) or fn.get("is_deleted") is True:
                continue

            # Chuẩn hóa kiểu ngày tháng
            c_at = fn.get("created_at")
            if isinstance(c_at, datetime):
                c_at_str = c_at.isoformat()
            else:
                c_at_str = str(c_at) if c_at else datetime.now(timezone.utc).isoformat()

            u_at = fn.get("updated_at")
            if isinstance(u_at, datetime):
                u_at_str = u_at.isoformat()
            else:
                u_at_str = str(u_at) if u_at else c_at_str

            raw_cat = fn.get("category") or "Ghi chú"
            if raw_cat not in ["Tiềm năng", "Lưu ý", "Ghi chú"]:
                raw_cat = "Ghi chú"
            note_item = {
                "id": fn.get("id") or str(uuid.uuid4()),
                "school_id": s_id,
                "school_name": s_name,
                "trip_id": fn.get("trip_id"),
                "trip_name": fn.get("trip_name") or "Chuyến đi",
                "campaign_id": fn.get("campaign_id"),
                "campaign_name": fn.get("campaign_name") or "Chiến dịch tuyển sinh",
                "waypoint_id": fn.get("waypoint_id"),
                "category": raw_cat,
                "content": fn.get("content", "").strip(),
                "author_name": fn.get("author_name") or "Cán bộ tuyển sinh",
                "created_at": c_at_str,
                "updated_at": u_at_str,
                "is_deleted": False
            }
            all_notes.append(note_item)

    # Sắp xếp mới nhất lên đầu
    all_notes.sort(key=lambda x: str(x.get("created_at", "")), reverse=True)
    return all_notes


@router.get("")
async def get_all_notes(
    trip_id: Optional[str] = None,
    campaign_id: Optional[str] = None,
    school_id: Optional[str] = None,
    category: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 100,
    page: int = 1
):
    """Lấy danh sách ghi chú thực địa (được đọc trực tiếp từ hồ sơ các trường học)"""
    db = get_database()
    raw_notes = await _get_all_embedded_notes(db)

    # Thống kê tổng quan toàn hệ thống (không phụ thuộc bộ lọc)
    total_all_notes = len(raw_notes)
    distinct_schools = len(set(n["school_id"] for n in raw_notes if n.get("school_id")))
    distinct_trips = len(set(n["trip_id"] for n in raw_notes if n.get("trip_id")))

    # Áp dụng bộ lọc
    filtered = raw_notes
    if trip_id:
        filtered = [n for n in filtered if n.get("trip_id") == trip_id]
    if campaign_id:
        filtered = [n for n in filtered if n.get("campaign_id") == campaign_id]
    if school_id:
        filtered = [n for n in filtered if n.get("school_id") == school_id or n.get("waypoint_id") == school_id]
    if category:
        filtered = [n for n in filtered if n.get("category") == category]
    if search and search.strip():
        q = search.strip().lower()
        filtered = [
            n for n in filtered
            if q in n.get("content", "").lower()
            or q in n.get("school_name", "").lower()
            or q in n.get("trip_name", "").lower()
            or q in n.get("author_name", "").lower()
        ]

    try:
        page_val = max(1, int(page))
    except Exception:
        page_val = 1
    try:
        limit_val = max(1, int(limit))
    except Exception:
        limit_val = 100

    skip = (page_val - 1) * limit_val
    paginated_items = filtered[skip : skip + limit_val]

    return {
        "items": paginated_items,
        "total": len(filtered),
        "page": page_val,
        "limit": limit_val,
        "stats": {
            "total_notes": total_all_notes,
            "total_schools": distinct_schools,
            "total_trips": distinct_trips
        }
    }


@router.post("")
async def create_note(data: CreateNoteRequest):
    """Tạo ghi chú mới và lưu trực tiếp vào mảng 'field_notes' của trường học tương ứng"""
    db = get_database()
    now_iso = datetime.now(timezone.utc).isoformat()
    note_id = str(uuid.uuid4())

    trip_name = data.trip_name
    campaign_name = data.campaign_name
    campaign_id = data.campaign_id

    # Tìm thông tin chuyến đi / chiến dịch nếu có
    if data.trip_id:
        trip = await db.admission_trips.find_one({"id": data.trip_id})
        if trip:
            trip_name = trip_name or trip.get("name")
            campaign_id = campaign_id or trip.get("campaign_id")

    if campaign_id:
        camp = await db.campaigns.find_one({"$or": [{"id": campaign_id}, {"_id": campaign_id}]})
        if camp:
            campaign_name = campaign_name or camp.get("name")

    # Tìm trường học mục tiêu trong collection 'schools'
    target_school = None
    if data.school_id:
        target_school = await db.schools.find_one({
            "$or": [{"id": data.school_id}, {"code": data.school_id}],
            "is_deleted": {"$ne": True}
        })
    if not target_school and data.school_name:
        target_school = await db.schools.find_one({
            "name": data.school_name.strip(),
            "is_deleted": {"$ne": True}
        })

    s_id = (target_school.get("id") or str(target_school["_id"])) if target_school else data.school_id
    s_name = target_school.get("name") if target_school else (data.school_name or "Điểm trường")

    note_doc = {
        "id": note_id,
        "trip_id": data.trip_id,
        "trip_name": trip_name or "Chuyến đi",
        "campaign_id": campaign_id,
        "campaign_name": campaign_name or "Chiến dịch tuyển sinh",
        "school_id": s_id,
        "school_name": s_name,
        "waypoint_id": data.waypoint_id,
        "category": data.category or "Ghi chú",
        "content": data.content.strip(),
        "author_name": data.author_name or "Cán bộ tuyển sinh",
        "created_at": now_iso,
        "updated_at": now_iso,
        "is_deleted": False
    }

    # Lưu trực tiếp vào mảng 'field_notes' của trường
    if target_school:
        await db.schools.update_one(
            {"_id": target_school["_id"]},
            {
                "$push": {"field_notes": note_doc},
                "$set": {
                    "notes": data.content.strip(),
                    "updated_at": datetime.now(timezone.utc)
                }
            }
        )
    else:
        # Nếu chưa có hồ sơ trường học, tạo mới hồ sơ trường tối giản
        new_school = {
            "id": s_id or str(uuid.uuid4()),
            "name": s_name,
            "address": "Chưa có địa chỉ",
            "lat": 10.7769,
            "lng": 106.7009,
            "notes": data.content.strip(),
            "field_notes": [note_doc],
            "created_at": datetime.now(timezone.utc),
            "updated_at": datetime.now(timezone.utc),
            "is_deleted": False
        }
        await db.schools.insert_one(new_school)

    # Đồng bộ vào waypoint của chuyến đi nếu có
    if data.trip_id and (data.waypoint_id or data.school_id):
        wp_query = {
            "trip_id": data.trip_id,
            "$or": [
                {"id": data.waypoint_id},
                {"school_id": data.school_id},
                {"id": data.school_id}
            ]
        }
        await db.campaign_waypoints.update_many(
            wp_query,
            {"$set": {"notes": data.content.strip()}}
        )

    return note_doc


@router.put("/{note_id}")
async def update_note(note_id: str, data: UpdateNoteRequest):
    """Cập nhật nội dung ghi chú bên trong trường học"""
    db = get_database()
    now_iso = datetime.now(timezone.utc).isoformat()

    school = await db.schools.find_one({"field_notes.id": note_id})
    if not school:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy ghi chú")

    set_fields = {"field_notes.$.updated_at": now_iso, "updated_at": datetime.now(timezone.utc)}
    if data.content is not None:
        set_fields["field_notes.$.content"] = data.content.strip()
        set_fields["notes"] = data.content.strip()
    if data.category is not None:
        set_fields["field_notes.$.category"] = data.category
    if data.school_name is not None:
        set_fields["field_notes.$.school_name"] = data.school_name

    await db.schools.update_one(
        {"field_notes.id": note_id},
        {"$set": set_fields}
    )

    # Lấy lại note vừa cập nhật
    updated_school = await db.schools.find_one({"field_notes.id": note_id})
    for fn in (updated_school.get("field_notes") or []):
        if fn.get("id") == note_id:
            fn["school_id"] = updated_school.get("id") or str(updated_school["_id"])
            fn["school_name"] = updated_school.get("name")
            return fn

    return {"message": "Đã cập nhật ghi chú", "id": note_id}


@router.delete("/{note_id}")
async def delete_note(note_id: str):
    """Xóa ghi chú (soft delete trong trường học)"""
    db = get_database()
    now_iso = datetime.now(timezone.utc).isoformat()

    res = await db.schools.update_one(
        {"field_notes.id": note_id},
        {
            "$set": {
                "field_notes.$.is_deleted": True,
                "field_notes.$.updated_at": now_iso
            }
        }
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy ghi chú")
    return {"message": "Đã xóa ghi chú thành công", "id": note_id}


@router.get("/export/excel")
async def export_notes_excel(
    trip_id: Optional[str] = None,
    campaign_id: Optional[str] = None,
    school_id: Optional[str] = None,
    category: Optional[str] = None,
    search: Optional[str] = None
):
    """Xuất toàn bộ ghi chú thực địa từ các trường học ra file Excel"""
    db = get_database()
    raw_notes = await _get_all_embedded_notes(db)

    filtered = raw_notes
    if trip_id:
        filtered = [n for n in filtered if n.get("trip_id") == trip_id]
    if campaign_id:
        filtered = [n for n in filtered if n.get("campaign_id") == campaign_id]
    if school_id:
        filtered = [n for n in filtered if n.get("school_id") == school_id or n.get("waypoint_id") == school_id]
    if category:
        filtered = [n for n in filtered if n.get("category") == category]
    if search and search.strip():
        q = search.strip().lower()
        filtered = [
            n for n in filtered
            if q in n.get("content", "").lower()
            or q in n.get("school_name", "").lower()
            or q in n.get("trip_name", "").lower()
            or q in n.get("author_name", "").lower()
        ]

    notes = filtered

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ghi chú thực địa"

    font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Tiêu đề bảng
    ws.merge_cells("A1:H1")
    ws["A1"] = "DANH SÁCH GHI CHÚ VÀ ĐÁNH GIÁ THỰC ĐỊA TUYỂN SINH"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    now_vn = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Thời điểm xuất file: {now_vn} | Tổng số ghi chú: {len(notes)}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    headers = [
        "STT",
        "Thời Gian Ghi",
        "Điểm Trường",
        "Phân Loại / Tag",
        "Chuyến Đi / Đợt",
        "Chiến Dịch",
        "Cán Bộ Ghi",
        "Nội Dung Ghi Chú & Đánh Giá"
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for idx, n in enumerate(notes, 1):
        r_idx = 4 + idx
        ws.row_dimensions[r_idx].height = 24

        c_time = n.get("created_at")
        if isinstance(c_time, datetime):
            c_time_str = c_time.strftime("%d/%m/%Y %H:%M")
        elif isinstance(c_time, str):
            try:
                c_time_str = datetime.fromisoformat(c_time).strftime("%d/%m/%Y %H:%M")
            except Exception:
                c_time_str = c_time
        else:
            c_time_str = ""

        row_vals = [
            idx,
            c_time_str,
            n.get("school_name") or "Chưa rõ",
            n.get("category") or "Ghi chú",
            n.get("trip_name") or "Chuyến đi",
            n.get("campaign_name") or "Chiến dịch tuyển sinh",
            n.get("author_name") or "Cán bộ tuyển sinh",
            n.get("content") or ""
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_bold if col_idx in (1, 3) else font_normal
            cell.border = thin_border
            align_h = "center" if col_idx in (1, 2, 4) else ("left" if col_idx in (3, 5, 6, 8) else "center")
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=(col_idx == 8))
            if idx % 2 == 0:
                cell.fill = fill_even

    # Căn chỉnh độ rộng cột
    col_widths = [8, 18, 30, 20, 28, 25, 20, 50]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Danh_sach_ghi_chu_tuyen_sinh_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
